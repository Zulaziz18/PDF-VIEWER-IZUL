#!/usr/bin/env python3
"""The Word-conversion proof (7.1.0): a .docx made by the application opens in
readers that are not the application, and says what the PDF said.

For every case of `make_cases.py`:

1. The application's conversion runs on it (`izul-bench --bin docx-proof`:
   `Document::page_layout` for every page, then `izul_docx::analyse` and
   `izul_docx::build` — the routines the export runs).
2. python-docx opens the package (a structural check: parts, relationships,
   content types) and reads its paragraphs; every picture in it decodes with
   Pillow.
3. LibreOffice Writer opens it and saves it as PDF — a second, independent
   reader that has to lay the document out, not only parse it.
4. The words: the source's words (poppler `pdftotext`, which shares no code
   with PDFium) must be found again, in order, in the .docx and in
   LibreOffice's PDF. Measured as the share of source words in the longest
   common ordered matching (difflib), 0.97 or more. The sideways margin line
   is left out of the source on purpose — the margin strip is cropped away —
   because the conversion skips text that is not upright, and says so.
5. The look, where it is the point of the case: the letter's date is right
   aligned and its subject centred, the article's title is 17 pt bold and
   centred, its authors italic, the photo and the figure are there, the page
   count of a one-column document is unchanged.
6. Mutations that must be caught: a .docx with every eighth paragraph taken
   out, and one with a picture's bytes cut short. A check that passes them
   is a check that cannot fail.

    python3 tools/docx-proof/run.py [--out bench/results/docx-proof.txt]

Needs poppler-utils, LibreOffice Writer (`libreoffice-writer`), python-docx
and Pillow. Exits non-zero if anything fails.
"""
import argparse
import difflib
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
THRESHOLD = 0.97
SIDEWAYS = "Diunduh dari contoh.invalid pada 1 Oktober 2026"

try:
    import docx  # python-docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from PIL import Image
except ImportError as e:  # pragma: no cover
    sys.exit(f"butuh python-docx dan Pillow: {e}")


def words(text):
    return re.findall(r"[0-9A-Za-zÀ-ÿ]+", text.lower())


def recall(source, got):
    if not source:
        return 1.0
    m = difflib.SequenceMatcher(None, source, got, autojunk=False)
    return sum(b.size for b in m.get_matching_blocks()) / len(source)


def pdftotext(path, skip_margin=False):
    """`skip_margin` leaves out the left margin strip, where the sideways line
    is: pdftotext breaks that line up and scatters it, so it cannot be taken
    out of the text afterwards."""
    crop = ["-x", "45", "-y", "0", "-W", "2000", "-H", "2000"] if skip_margin else []
    return subprocess.run(["pdftotext", "-raw", *crop, str(path), "-"], check=True, capture_output=True, text=True).stdout


def page_count(path):
    out = subprocess.run(["pdfinfo", str(path)], check=True, capture_output=True, text=True).stdout
    return int(re.search(r"Pages:\s+(\d+)", out).group(1))


def docx_text(path):
    d = docx.Document(str(path))
    return "\n".join(p.text for p in d.paragraphs), d


def pictures_ok(path):
    """Every picture part decodes. Returns the count."""
    n = 0
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if name.startswith("word/media/"):
                Image.open(io.BytesIO(z.read(name))).load()
                n += 1
    return n


def libreoffice_pdf(path, work):
    home = work / "lo-home"
    home.mkdir(exist_ok=True)
    env = dict(os.environ, HOME=str(home))
    subprocess.run(
        ["soffice", "--headless", "--norestore", "--convert-to", "pdf", str(path), "--outdir", str(work)],
        check=True,
        capture_output=True,
        env=env,
        timeout=300,
    )
    out = work / (path.stem + ".pdf")
    if not out.exists():
        raise RuntimeError("LibreOffice tidak menghasilkan PDF")
    return out


def convert(src, out):
    exe = ROOT / "target" / "debug" / ("docx-proof.exe" if os.name == "nt" else "docx-proof")
    r = subprocess.run([str(exe), str(src), str(out)], check=True, capture_output=True, text=True)
    return json.loads(r.stdout.strip().splitlines()[-1])


def paragraph(d, starts):
    for p in d.paragraphs:
        if p.text.startswith(starts):
            return p
    return None


def look_checks(name, d, src_pages, lo_pages):
    """The case-specific look. Returns a list of (what, ok)."""
    out = []
    if name == "surat":
        date = paragraph(d, "Kota Contoh")
        subject = paragraph(d, "Perihal")
        out.append(("tanggal rata kanan", date is not None and date.alignment == WD_ALIGN_PARAGRAPH.RIGHT))
        out.append(("perihal rata tengah", subject is not None and subject.alignment == WD_ALIGN_PARAGRAPH.CENTER))
        out.append(("perihal tebal", subject is not None and all(r.bold for r in subject.runs if r.text.strip())))
        out.append(("jumlah halaman sama", lo_pages == src_pages))
    if name == "jurnal":
        title = paragraph(d, "Waktu Tunggu")
        authors = paragraph(d, "Rina Contoh")
        out.append(("judul 17 pt", title is not None and all(r.font.size and abs(r.font.size.pt - 17) < 0.6 for r in title.runs if r.text.strip())))
        out.append(("judul tebal", title is not None and all(r.bold for r in title.runs if r.text.strip())))
        out.append(("judul rata tengah", title is not None and title.alignment == WD_ALIGN_PARAGRAPH.CENTER))
        out.append(("penulis miring", authors is not None and all(r.italic for r in authors.runs if r.text.strip())))
        out.append(("baris tabel bertab", any("\t" in p.text and p.text.startswith("Konsultasi") for p in d.paragraphs)))
    if name in ("foto", "pindaian", "surat"):
        out.append(("jumlah halaman sama", lo_pages == src_pages) if name != "surat" else ("-", True))
    return [x for x in out if x[0] != "-"]


def mutate_drop_paragraphs(src, dst):
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                # Whole <w:p> elements (they never nest), every eighth one.
                count = [0]

                def drop(m):
                    count[0] += 1
                    return "" if count[0] % 8 == 0 else m.group(0)

                data = re.sub(r"<w:p>.*?</w:p>", drop, data.decode()).encode()
            zout.writestr(item, data)


def mutate_cut_picture(src, dst):
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("word/media/"):
                data = data[: len(data) // 3]
            zout.writestr(item, data)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=None)
    args = ap.parse_args()
    for tool in ("pdftotext", "pdfinfo", "soffice"):
        if not shutil.which(tool):
            sys.exit(f"{tool} tidak ada (poppler-utils, libreoffice-writer)")
    subprocess.run(["cargo", "build", "-q", "-p", "izul-bench", "--bin", "docx-proof"], check=True, cwd=ROOT)

    lines = []
    failed = False
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        cases = work / "cases"
        subprocess.run([sys.executable, str(ROOT / "tools/docx-proof/make_cases.py"), str(cases)], check=True, capture_output=True)
        expected_pictures = {"jurnal": 1, "surat": 0, "foto": 1, "pindaian": 1}
        for name in ("jurnal", "surat", "foto", "pindaian"):
            src = cases / f"{name}.pdf"
            out = work / f"{name}.docx"
            stats = convert(src, out)
            source_words = words(pdftotext(src, skip_margin=name == "jurnal"))
            text, d = docx_text(out)
            pics = pictures_ok(out)
            lo = libreoffice_pdf(out, work)
            lo_words = words(pdftotext(lo))
            r_docx = recall(source_words, words(text))
            r_lo = recall(source_words, lo_words)
            src_pages, lo_pages = page_count(src), page_count(lo)
            checks = [
                ("python-docx membuka", True),
                (f"gambar {pics}/{expected_pictures[name]} terbaca", pics == expected_pictures[name]),
                (f"kata di .docx {r_docx:.3f}", r_docx >= THRESHOLD),
                (f"kata di PDF LibreOffice {r_lo:.3f}", r_lo >= THRESHOLD),
            ] + look_checks(name, d, src_pages, lo_pages)
            ok = all(c[1] for c in checks)
            failed |= not ok
            lines.append(
                f"{'LULUS' if ok else 'GAGAL'}  {name:9} {src_pages}→{lo_pages} hlm  "
                f"{stats['paragraphs']} paragraf  {stats['pictures']} gambar  "
                f"{len(source_words)} kata  {stats['bytes']} B  {stats['ms']} ms"
            )
            for what, good in checks:
                lines.append(f"         {'ok ' if good else 'XX '} {what}")
            if name == "jurnal":
                lines.append(f"         (teks miring dilewati: {stats['skipped_turned']} huruf — baris tepi '{SIDEWAYS}')")
            if name == "pindaian":
                lines.append(f"         (halaman tanpa teks: {stats['pages_without_text']} — masuk sebagai gambar)")

        # Mutations: each must be caught.
        lines.append("")
        lines.append("Mutasi yang wajib tertangkap:")
        jurnal = work / "jurnal.docx"
        src_words = words(pdftotext(cases / "jurnal.pdf", skip_margin=True))
        dropped = work / "mutasi-paragraf.docx"
        mutate_drop_paragraphs(jurnal, dropped)
        r = recall(src_words, words(docx_text(dropped)[0]))
        caught = r < THRESHOLD
        failed |= not caught
        lines.append(f"  {'tertangkap' if caught else 'LOLOS!'}  satu dari delapan paragraf dibuang: kata {r:.3f}")
        cut = work / "mutasi-gambar.docx"
        mutate_cut_picture(work / "foto.docx", cut)
        try:
            pictures_ok(cut)
            caught = False
        except Exception:
            caught = True
        failed |= not caught
        lines.append(f"  {'tertangkap' if caught else 'LOLOS!'}  byte gambar dipotong")

    lo_version = subprocess.run(["soffice", "--version"], capture_output=True, text=True).stdout.strip()
    report = "\n".join(
        [
            "Bukti konversi PDF ke Word (7.1.0) — tools/docx-proof/run.py",
            f"Pembaca: python-docx {docx.__version__ if hasattr(docx, '__version__') else ''}, {lo_version}, poppler pdftotext",
            f"Ambang kata berurutan: {THRESHOLD}",
            "",
            *lines,
            "",
            "HASIL: " + ("GAGAL" if failed else "LULUS"),
        ]
    )
    print(report)
    if args.out:
        Path(args.out).write_text(report + "\n", encoding="utf-8")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
